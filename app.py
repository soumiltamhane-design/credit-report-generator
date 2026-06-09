import streamlit as st
import json, re, io, os, base64, requests
import google.generativeai as genai
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

st.set_page_config(page_title="Credit Report Generator", page_icon="📋", layout="wide", initial_sidebar_state="expanded")

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Serif+Display:ital@0;1&family=DM+Sans:wght@300;400;500;600&display=swap');
html, body, [class*="css"] { font-family: 'DM Sans', sans-serif; }
.stApp { background: #0f0f13; color: #e8e4dc; }
[data-testid="stSidebar"] { background: #16161c; border-right: 1px solid #2a2a35; }
h1 { font-family: 'DM Serif Display', serif !important; font-size: 2.2rem !important; color: #e8e4dc !important; letter-spacing: -1px !important; }
h2, h3 { font-family: 'DM Serif Display', serif !important; color: #e8e4dc !important; }
label, .stTextInput label, .stSelectbox label, .stTextArea label { color: #9b9688 !important; font-size: 0.75rem !important; font-weight: 600 !important; letter-spacing: 1px !important; text-transform: uppercase !important; }
.stTextInput input, .stTextArea textarea { background: #1c1c24 !important; border: 1px solid #2a2a35 !important; border-radius: 8px !important; color: #e8e4dc !important; }
.stTextInput input:focus, .stTextArea textarea:focus { border-color: #c9a84c !important; box-shadow: 0 0 0 2px rgba(201,168,76,0.15) !important; }
.stButton button { background: #c9a84c !important; color: #0f0f13 !important; border: none !important; border-radius: 8px !important; font-weight: 600 !important; padding: 0.6rem 1.5rem !important; }
.stButton button:hover { background: #e0bc62 !important; }
.stDownloadButton button { background: #1c1c24 !important; color: #c9a84c !important; border: 1px solid #c9a84c !important; border-radius: 8px !important; font-weight: 600 !important; }
.stDownloadButton button:hover { background: #c9a84c !important; color: #0f0f13 !important; }
.stRadio label { color: #e8e4dc !important; text-transform: none !important; letter-spacing: 0 !important; font-size: 0.9rem !important; font-weight: 400 !important; }
hr { border-color: #2a2a35 !important; margin: 20px 0 !important; }
#MainMenu {visibility: hidden;} footer {visibility: hidden;} header {visibility: hidden;}
.url-section { background: #16161c; border: 1px solid #2a2a35; border-radius: 10px; padding: 16px; margin-bottom: 12px; }
.period-badge { display: inline-block; background: rgba(201,168,76,0.15); color: #c9a84c; border: 1px solid rgba(201,168,76,0.3); border-radius: 4px; padding: 2px 8px; font-size: 0.7rem; font-weight: 700; letter-spacing: 1px; }
</style>
""", unsafe_allow_html=True)

# ── DOCX builder ──────────────────────────────────────────────
DARK_NAVY=RGBColor(0x1A,0x1A,0x2E); BLUE=RGBColor(0x25,0x63,0xEB); WHITE=RGBColor(0xFF,0xFF,0xFF)
DARK_TEXT=RGBColor(0x1F,0x29,0x37); MUTED_C=RGBColor(0x47,0x55,0x69)

def _bg(cell,h):
    s=OxmlElement('w:shd'); s.set(qn('w:val'),'clear'); s.set(qn('w:color'),'auto'); s.set(qn('w:fill'),h)
    cell._tc.get_or_add_tcPr().append(s)

def _borders(cell):
    b=OxmlElement('w:tcBorders')
    for side in ('top','bottom','left','right'):
        e=OxmlElement(f'w:{side}'); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'2'); e.set(qn('w:space'),'0'); e.set(qn('w:color'),'CCCCCC'); b.append(e)
    cell._tc.get_or_add_tcPr().append(b)

def _cw(table,widths):
    tg=OxmlElement('w:tblGrid')
    for w in widths:
        gc=OxmlElement('w:gridCol'); gc.set(qn('w:w'),str(w)); tg.append(gc)
    table._tbl.insert(1,tg)
    for row in table.rows:
        for cell,w in zip(row.cells,widths):
            tw=OxmlElement('w:tcW'); tw.set(qn('w:w'),str(w)); tw.set(qn('w:type'),'dxa'); cell._tc.get_or_add_tcPr().append(tw)

def _r(para,text,bold=False,size=10.5,color=None):
    run=para.add_run(str(text)); run.bold=bold; run.font.name='Calibri'; run.font.size=Pt(size)
    if color: run.font.color.rgb=color
    return run

def _hd(doc,text):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(12); p.paragraph_format.space_after=Pt(4)
    _r(p,text.upper(),bold=True,size=9.5,color=BLUE)
    pBdr=OxmlElement('w:pBdr'); btm=OxmlElement('w:bottom')
    btm.set(qn('w:val'),'single'); btm.set(qn('w:sz'),'8'); btm.set(qn('w:space'),'1'); btm.set(qn('w:color'),'2563EB')
    pBdr.append(btm); p._p.get_or_add_pPr().append(pBdr)

def build_docx(data):
    doc=Document()
    for section in doc.sections:
        section.page_height=Cm(29.7); section.page_width=Cm(21.0)
        section.left_margin=Inches(1); section.right_margin=Inches(1)
        section.top_margin=Inches(1); section.bottom_margin=Inches(1)
    entity=data['entityType']
    periods=data.get('financialPeriods',['H1FY26','H1FY25','31.03.2025','31.03.2024'])
    for text,sz,fill,clr in [(data['issuerName'].upper(),18,'1A1A2E',WHITE),('Credit Review & Analysis',12,'1A1A2E',RGBColor(0x94,0xA3,0xB8)),(' ',4,'2563EB',WHITE)]:
        p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0)
        shd=OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),fill); p._p.get_or_add_pPr().append(shd)
        _r(p,text,bold=(sz==18),size=sz,color=clr)
    _hd(doc,'Company Profile')
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(8); _r(p,data.get('companyProfile',''),size=10.5,color=DARK_TEXT)
    _hd(doc,'Issuer Information')
    t=doc.add_table(rows=6,cols=2); t.style='Table Grid'; _cw(t,[2700,6300])
    for i,(label,value) in enumerate([('Issuer',data['issuerName']),('Industry as per NIC code',data.get('industry','')),('Sector',data.get('sector','')),('Review Period',data.get('reviewPeriod','')),('Prepared By',data['preparedBy']),('Reviewed By',data['reviewedBy'])]):
        lc,vc=t.rows[i].cells[0],t.rows[i].cells[1]
        _bg(lc,'F1F5F9'); _borders(lc); _bg(vc,'FFFFFF'); _borders(vc)
        for cell,txt,bold,clr in [(lc,label,True,MUTED_C),(vc,value,False,DARK_TEXT)]:
            p=cell.paragraphs[0]; p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(2); _r(p,txt,bold=bold,size=9.5,color=clr)
    doc.add_paragraph().paragraph_format.space_after=Pt(4)
    _hd(doc,'Brief Detail of Investment under Review')
    p2=doc.add_paragraph(); p2.paragraph_format.space_after=Pt(4); _r(p2,'(Rs. in crores)',size=9,color=MUTED_C)
    invs=[i for i in data.get('investments',[]) if i.get('security')]
    if invs:
        hdrs=['Security / Date of Maturity','Holding Yield','Credit Rating Agency','Credit Rating','Business House','Total Nominal Value (F.V)','SH','PH']
        ti=doc.add_table(rows=1+len(invs),cols=8); ti.style='Table Grid'
        ws=[2100,680,1100,700,1100,1100,680,680]; total=9360; s=sum(ws); ws=[int(w*total/s) for w in ws]; ws[-1]+=total-sum(ws); _cw(ti,ws)
        for ci,h in enumerate(hdrs):
            cell=ti.rows[0].cells[ci]; _bg(cell,'1A1A2E'); _borders(cell)
            p=cell.paragraphs[0]; p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(2); _r(p,h,bold=True,size=8,color=WHITE)
        for ri,inv in enumerate(invs):
            row=ti.rows[ri+1]; bgc='FFFFFF' if ri%2==0 else 'F8FAFC'
            for ci,v in enumerate([inv.get('security',''),inv.get('yield',''),inv.get('agency',''),inv.get('rating',''),inv.get('businessHouse',''),inv.get('fv',''),inv.get('sh',''),inv.get('ph','')]):
                cell=row.cells[ci]; _bg(cell,bgc); _borders(cell)
                p=cell.paragraphs[0]; p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(2); _r(p,v or '',size=9,color=DARK_TEXT)
    doc.add_paragraph().paragraph_format.space_after=Pt(4)
    _hd(doc,'Financial Strength')
    p3=doc.add_paragraph(); p3.paragraph_format.space_after=Pt(4); _r(p3,'(Rs. in crores)',size=9,color=MUTED_C)
    fin=list(data.get('financialData',[]))
    for m in data.get('conditionalMetrics',[]):
        if entity in m.get('includeFor',[]): fin.append(m)
    if fin:
        tf=doc.add_table(rows=1+len(fin),cols=1+len(periods)); tf.style='Table Grid'
        lw=2500; dw=(9360-lw)//len(periods); ws=[lw]+[dw]*len(periods); ws[-1]+=9360-sum(ws); _cw(tf,ws)
        for ci,h in enumerate(['Particulars']+list(periods)):
            cell=tf.rows[0].cells[ci]; _bg(cell,'1A1A2E'); _borders(cell)
            p=cell.paragraphs[0]; p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(2); _r(p,h,bold=True,size=9,color=WHITE)
        for ri,m in enumerate(fin):
            row=tf.rows[ri+1]; bgc='FFFFFF' if ri%2==0 else 'EFF6FF'
            for ci,v in enumerate([m.get('metric','')]+[str(x) if x else '—' for x in m.get('values',[])]):
                cell=row.cells[ci]; _bg(cell,bgc); _borders(cell)
                p=cell.paragraphs[0]; p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(2); _r(p,v,bold=(ci==0),size=9.5,color=DARK_TEXT)
    doc.add_paragraph().paragraph_format.space_after=Pt(4)
    _hd(doc,'Comments')
    for c in data.get('comments',[]):
        p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(4)
        _r(p,c['heading']+': ',bold=True,size=10.5,color=BLUE); _r(p,c['text'],size=10.5,color=DARK_TEXT)
    doc.add_paragraph().paragraph_format.space_after=Pt(4)
    _hd(doc,'Recommendation')
    p4=doc.add_paragraph(); p4.paragraph_format.space_after=Pt(16); _r(p4,data.get('recommendation',''),size=10.5,color=DARK_TEXT)
    t2=doc.add_table(rows=1,cols=2); t2.style='Table Grid'; _cw(t2,[4680,4680])
    for ci,(role,name) in enumerate([('Fund Manager',data['preparedBy']),('CIO',data['reviewedBy'])]):
        cell=t2.rows[0].cells[ci]; _bg(cell,'F1F5F9'); _borders(cell)
        p=cell.paragraphs[0]; p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(4)
        _r(p,role+': ',bold=True,size=10,color=MUTED_C); _r(p,name,size=10,color=DARK_TEXT)
    buf=io.BytesIO(); doc.save(buf); buf.seek(0); return buf

# ── PDF extraction ────────────────────────────────────────────
def extract_from_pdf_bytes(pdf_bytes, label):
    text = ''
    try:
        import fitz
        doc = fitz.open(stream=pdf_bytes, filetype='pdf')
        page_keywords = ['profit','loss','income','balance sheet','assets','npa','gnpa','capital adequacy','crar','net interest','borrowings','advances','deposits','net worth','stage 3','ratios','notes to']
        scored = []
        for i in range(len(doc)):
            page = doc[i]; t = page.get_text(); tl = t.lower()
            score = sum(1 for kw in page_keywords if kw in tl)
            if score >= 2: scored.append((score, i, t))
        scored.sort(reverse=True)
        for score, pnum, t in scored[:15]:
            tl = t.lower()
            if any(kw in tl for kw in ['profit and loss','statement of profit']): lbl='[P&L]'
            elif any(kw in tl for kw in ['balance sheet','financial position']): lbl='[BALANCE SHEET]'
            elif any(kw in tl for kw in ['capital adequacy','crar','tier i']): lbl='[NOTES-CAPITAL]'
            elif any(kw in tl for kw in ['gross npa','net npa','stage 3','asset quality']): lbl='[NOTES-NPA]'
            elif any(kw in tl for kw in ['borrowings','debentures','ncd']): lbl='[NOTES-BORROWINGS]'
            elif any(kw in tl for kw in ['ratio','nim','roa','roe','highlights']): lbl='[KEY RATIOS]'
            else: lbl='[FINANCIAL PAGE]'
            text += f'\n\n{"="*40}\n{lbl} pg{pnum+1}\n{"="*40}\n{t}'
        doc.close()
    except Exception as e:
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
                text = '\n'.join([p.extract_text() or '' for p in pdf.pages[:30]])
        except: pass
    return f'\n\n{"#"*40}\n[{label}]\n{"#"*40}\n{text[:5000]}'

def fetch_pdf(url, label):
    if not url.strip(): return ''
    try:
        headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'}
        r = requests.get(url.strip(), headers=headers, timeout=60)
        ct = r.headers.get('Content-Type','').lower()
        if 'pdf' in ct or url.lower().endswith('.pdf') or '.pdf?' in url.lower():
            return extract_from_pdf_bytes(r.content, label)
        else:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(r.text,'html.parser')
            candidates = []
            base = url.split('/')[0]+'//'+url.split('/')[2]
            for tag in soup.find_all('a', href=True):
                href = tag['href']
                full = href if href.startswith('http') else base+href
                if '.pdf' in full.lower():
                    txt = tag.get_text(strip=True).lower()
                    score = sum(3 for kw in ['annual report','quarterly result','financial result','q1','q2','h1','fy25','fy26'] if kw in txt+full.lower())
                    score -= sum(5 for kw in ['csr','agm','notice','sustainability'] if kw in txt+full.lower())
                    candidates.append((score, full))
            if candidates:
                candidates.sort(reverse=True)
                best = candidates[0][1]
                r2 = requests.get(best, headers=headers, timeout=60)
                return extract_from_pdf_bytes(r2.content, label)
    except Exception as e:
        st.warning(f'Could not fetch {label}: {e}')
    return ''

# ── Sidebar ───────────────────────────────────────────────────
with st.sidebar:
    st.markdown("# 📋 Credit Report\nGenerator")
    st.markdown("---")
    st.markdown('<p style="color:#9b9688;font-size:0.8rem">Health Insurance Premium Corpus</p>', unsafe_allow_html=True)
    st.markdown("### Configuration")
    api_key = st.text_input("Gemini API Key", type="password", placeholder="AIza...")
    st.markdown("---")
    st.markdown("### Entity Type")
    entity_type = st.radio("", ["Bank / SFB", "NBFC / IFC", "PSU Finance", "Apex Institution"], label_visibility="collapsed")
    entity_map = {"Bank / SFB":"bank","NBFC / IFC":"nbfc","PSU Finance":"psu","Apex Institution":"apex"}
    entity_code = entity_map[entity_type]
    st.markdown("---")
    st.markdown('<p style="color:#9b9688;font-size:0.75rem">Cost per report: ~₹0 | Free tier: 50/day</p>', unsafe_allow_html=True)

# ── Main ──────────────────────────────────────────────────────
st.markdown("# Credit Report Generator")
st.markdown('<p style="color:#9b9688;margin-top:-12px;margin-bottom:28px">AI-powered credit analysis for debt investment decisions</p>', unsafe_allow_html=True)

col1, col2 = st.columns([1.2, 1], gap="large")

with col1:
    st.markdown("### Issuer Details")
    c1,c2 = st.columns(2)
    with c1: issuer_name = st.text_input("Issuer Name", placeholder="Power Finance Corporation Ltd")
    with c2: review_period = st.text_input("Review Period", value="Year period ended September 30, 2025")
    c3,c4 = st.columns(2)
    with c3: prepared_by = st.text_input("Prepared By", placeholder="Analyst Name")
    with c4: reviewed_by = st.text_input("Reviewed By", placeholder="CIO Name")
    c5,c6 = st.columns(2)
    with c5: industry = st.text_input("Industry", value="Financial and insurance activities")
    with c6: sector = st.text_input("Sector", value="Finance")

    st.markdown("---")
    st.markdown("### Source Documents")
    st.markdown('<p style="color:#9b9688;font-size:0.85rem">Paste direct PDF links from BSE/NSE/company IR page. H1 figures are auto-calculated from Q1+Q2.</p>', unsafe_allow_html=True)

    tab1, tab2 = st.tabs(["🔗 Paste URLs", "📁 Upload Files"])

    with tab1:
        st.markdown('<p style="color:#c9a84c;font-size:0.75rem;font-weight:700;letter-spacing:1px">H1FY26 = Q1FY26 + Q2FY26 &nbsp;|&nbsp; H1FY25 = Q1FY25 + Q2FY25</p>', unsafe_allow_html=True)
        q1fy26_url = st.text_input("Q1FY26 — Apr-Jun 2025 Quarterly Results", placeholder="https://...pdf")
        q2fy26_url = st.text_input("Q2FY26 — Jul-Sep 2025 Quarterly Results", placeholder="https://...pdf")
        q1fy25_url = st.text_input("Q1FY25 — Apr-Jun 2024 Quarterly Results", placeholder="https://...pdf")
        q2fy25_url = st.text_input("Q2FY25 — Jul-Sep 2024 Quarterly Results", placeholder="https://...pdf")
        fy25_url   = st.text_input("FY25 Annual Report — 31.03.2025", placeholder="https://...pdf")
        fy24_url   = st.text_input("FY24 Annual Report — 31.03.2024", placeholder="https://...pdf")

    with tab2:
        st.markdown('<p style="color:#9b9688;font-size:0.8rem">Upload if you have PDFs saved locally. Label each file clearly.</p>', unsafe_allow_html=True)
        uploaded_files = st.file_uploader("Upload PDFs (select multiple)", type=['pdf'], accept_multiple_files=True,
                                           help="Hold Ctrl to select multiple files")
        if uploaded_files:
            st.markdown(f'<p style="color:#c9a84c;font-size:0.8rem">✅ {len(uploaded_files)} file(s) uploaded</p>', unsafe_allow_html=True)
            for f in uploaded_files:
                st.markdown(f'<p style="color:#9b9688;font-size:0.75rem">• {f.name}</p>', unsafe_allow_html=True)

    st.markdown("---")
    st.markdown("### Investments under Review")
    st.markdown('<p style="color:#9b9688;font-size:0.8rem">Add one block per bond</p>', unsafe_allow_html=True)

    if 'investments' not in st.session_state:
        st.session_state.investments = [{'security':'','yield':'','agency':'','rating':'','businessHouse':'','fv':'','sh':'','ph':''}]

    investments_data = []
    for i, inv in enumerate(st.session_state.investments):
        with st.expander(f"Bond {i+1}" + (f" — {inv['security']}" if inv['security'] else ""), expanded=(i==0)):
            r1c1,r1c2 = st.columns([2,1])
            with r1c1: sec = st.text_input("Security / Maturity", value=inv['security'], placeholder="8.03% NCD May 02, 2026", key=f"sec_{i}")
            with r1c2: yld = st.text_input("Holding Yield %", value=inv['yield'], placeholder="7.55", key=f"yld_{i}")
            r2c1,r2c2,r2c3 = st.columns(3)
            with r2c1: agency = st.text_input("Rating Agency", value=inv['agency'], placeholder="CRISIL", key=f"agency_{i}")
            with r2c2: rating = st.text_input("Rating", value=inv['rating'], placeholder="AAA", key=f"rating_{i}")
            with r2c3: bh = st.text_input("Business House", value=inv['businessHouse'], placeholder="PFC Group", key=f"bh_{i}")
            r3c1,r3c2,r3c3 = st.columns(3)
            with r3c1: fv = st.text_input("Face Value (Cr)", value=inv['fv'], placeholder="25", key=f"fv_{i}")
            with r3c2: sh = st.text_input("SH", value=inv['sh'], placeholder="—", key=f"sh_{i}")
            with r3c3: ph = st.text_input("PH", value=inv['ph'], placeholder="25", key=f"ph_{i}")
            investments_data.append({'security':sec,'yield':yld,'agency':agency,'rating':rating,'businessHouse':bh,'fv':fv,'sh':sh,'ph':ph})

    if st.button("＋ Add another bond"):
        st.session_state.investments.append({'security':'','yield':'','agency':'','rating':'','businessHouse':'','fv':'','sh':'','ph':''})
        st.rerun()

with col2:
    st.markdown("### Generate Report")
    s1,s2,s3 = st.columns(3)
    with s1: st.markdown(f'<div style="background:#1c1c24;border:1px solid #2a2a35;border-radius:10px;padding:16px;text-align:center"><div style="font-family:serif;font-size:1.8rem;color:#c9a84c">~60s</div><div style="font-size:0.7rem;color:#9b9688;text-transform:uppercase;letter-spacing:1px">Per Report</div></div>', unsafe_allow_html=True)
    with s2: st.markdown(f'<div style="background:#1c1c24;border:1px solid #2a2a35;border-radius:10px;padding:16px;text-align:center"><div style="font-family:serif;font-size:1.8rem;color:#c9a84c">₹0</div><div style="font-size:0.7rem;color:#9b9688;text-transform:uppercase;letter-spacing:1px">Cost</div></div>', unsafe_allow_html=True)
    with s3: st.markdown(f'<div style="background:#1c1c24;border:1px solid #2a2a35;border-radius:10px;padding:16px;text-align:center"><div style="font-family:serif;font-size:1.4rem;color:#c9a84c">{entity_type.split("/")[0].strip()}</div><div style="font-size:0.7rem;color:#9b9688;text-transform:uppercase;letter-spacing:1px">Entity</div></div>', unsafe_allow_html=True)

    st.markdown("")
    generate = st.button("✦ Generate Credit Report", use_container_width=True)

    if generate:
        errors = []
        if not api_key: errors.append("Gemini API key required")
        if not issuer_name: errors.append("Issuer name required")
        if not prepared_by: errors.append("Prepared By required")
        if not reviewed_by: errors.append("Reviewed By required")
        if errors:
            for e in errors: st.error(f"⚠️ {e}")
        else:
            entity_rules = {
                'bank': 'Bank/SFB. INCLUDE: NII, Deposits, CASA%, GNPA/Gross Stage 3%, NNPA/Net Stage 3%, CAR%, Tier I CAR%, ROA%, ROE%, NIM%. NO Debt/Equity.',
                'nbfc': 'NBFC/IFC. INCLUDE: NII, Loans & Advances, Borrowings, Net Worth, GNPA/Gross Stage 3%, NNPA/Net Stage 3%, CRAR%, Debt/Equity, EPS, NIM%. NO Deposits.',
                'psu' : 'PSU Finance. INCLUDE: NII, PAT, Net Worth, Loans, Borrowings, GNPA/Gross Stage 3%, NNPA/Net Stage 3%, CRAR%, Debt/Equity, EPS. NO Deposits.',
                'apex': 'Apex DFI. INCLUDE: NII, PAT, Net Worth, Advances, Borrowings, GNPA%, NNPA%, CRAR%, ROA%. Comment on GoI mandate.'
            }

            # ── Extract documents ─────────────────────────────
            extracted_text = ''
            doc_store = {'Q1FY26':'','Q2FY26':'','Q1FY25':'','Q2FY25':'','FY25':'','FY24':''}

            url_map = [('Q1FY26',q1fy26_url),('Q2FY26',q2fy26_url),('Q1FY25',q1fy25_url),('Q2FY25',q2fy25_url),('FY25',fy25_url),('FY24',fy24_url)]
            provided_urls = [(lbl,url) for lbl,url in url_map if url.strip()]

            if provided_urls:
                progress = st.progress(0, text="Fetching documents...")
                for idx,(lbl,url) in enumerate(provided_urls):
                    progress.progress((idx+1)/len(provided_urls), text=f"Fetching {lbl}...")
                    doc_store[lbl] = fetch_pdf(url, lbl)
                progress.empty()
                st.success(f"✅ {len(provided_urls)} document(s) fetched")

            # Handle uploaded files
            if 'uploaded_files' in dir() and uploaded_files:
                for uf in uploaded_files:
                    fname = uf.name.lower()
                    if 'q1' in fname and ('fy26' in fname or '2025' in fname and 'q1' in fname): lbl='Q1FY26'
                    elif 'q2' in fname and ('fy26' in fname or 'sep' in fname and '2025' in fname): lbl='Q2FY26'
                    elif 'q1' in fname and ('fy25' in fname or '2024' in fname): lbl='Q1FY25'
                    elif 'q2' in fname and ('fy25' in fname or 'sep' in fname and '2024' in fname): lbl='Q2FY25'
                    elif 'fy25' in fname or ('annual' in fname and '25' in fname): lbl='FY25'
                    elif 'fy24' in fname or ('annual' in fname and '24' in fname): lbl='FY24'
                    else: lbl=f'DOC_{uf.name[:10]}'
                    doc_store[lbl] = extract_from_pdf_bytes(uf.read(), lbl)
                st.success(f"✅ {len(uploaded_files)} file(s) processed")

            # Build combined text with H1 instructions
            parts = []
            if doc_store['Q1FY26'] or doc_store['Q2FY26']:
                parts.append(f"""{'='*50}
H1FY26 DATA (April 2025 - September 2025)
FLOW ITEMS (Total Income, NII, PAT): ADD Q1FY26 + Q2FY26
STOCK ITEMS (Assets, Loans, Borrowings, Net Worth): USE Q2FY26 value only
RATIO ITEMS (GNPA%, NNPA%, CRAR%, NIM%, ROE%): USE Q2FY26 value only
{'='*50}
Q1FY26 (Apr-Jun 2025):
{doc_store['Q1FY26'][:3000]}

Q2FY26 (Jul-Sep 2025):
{doc_store['Q2FY26'][:3000]}""")

            if doc_store['Q1FY25'] or doc_store['Q2FY25']:
                parts.append(f"""{'='*50}
H1FY25 DATA (April 2024 - September 2024)
FLOW ITEMS (Total Income, NII, PAT): ADD Q1FY25 + Q2FY25
STOCK ITEMS (Assets, Loans, Borrowings, Net Worth): USE Q2FY25 value only
RATIO ITEMS (GNPA%, NNPA%, CRAR%, NIM%, ROE%): USE Q2FY25 value only
{'='*50}
Q1FY25 (Apr-Jun 2024):
{doc_store['Q1FY25'][:3000]}

Q2FY25 (Jul-Sep 2024):
{doc_store['Q2FY25'][:3000]}""")

            if doc_store['FY25']:
                parts.append(f"{'='*50}\n31.03.2025 - FY25 ANNUAL REPORT (use directly)\n{'='*50}\n{doc_store['FY25'][:4000]}")
            if doc_store['FY24']:
                parts.append(f"{'='*50}\n31.03.2024 - FY24 ANNUAL REPORT (use directly)\n{'='*50}\n{doc_store['FY24'][:4000]}")

            extracted_text = '\n\n'.join(parts)[:14000] if parts else 'No documents provided. Use your knowledge of this company.'

            inv_text = '\n'.join(f"- {i.get('security')} | Yield {i.get('yield')}% | {i.get('rating')} ({i.get('agency')}) | FV Rs {i.get('fv')} Cr" for i in investments_data if i.get('security')) or 'No investments provided.'

            prompt = f"""You are a senior credit analyst at an Indian Health Insurance company (IRDAI-regulated).
Prepare a Credit Review & Analysis report.

ISSUER: {issuer_name}
ENTITY TYPE: {entity_code.upper()}
ENTITY RULES: {entity_rules[entity_code]}
INVESTMENTS: {inv_text}

{'='*55}
SOURCE FINANCIAL DATA:
{'='*55}
{extracted_text}

{'='*55}
CRITICAL INSTRUCTIONS:
{'='*55}

H1 CALCULATION — MANDATORY:
FLOW ITEMS (ADD Q1+Q2): Total Income, NII, PAT, Provisions, Fee Income
STOCK ITEMS (Q2 ONLY): Total Assets, Loans, Borrowings, Net Worth, Deposits
RATIO ITEMS (Q2 ONLY): GNPA%, NNPA%, CRAR%, NIM%, ROE%, ROA%
ANNUAL (USE DIRECTLY): 31.03.2025 from FY25, 31.03.2024 from FY24

FORMULAS (use when not stated):
- NII = Interest Income - Interest Expense
- GNPA% = Gross Stage 3 / Gross Loan Book x 100
- NNPA% = Net Stage 3 / Net Loan Book x 100
- NIM% = (H1 NII x 2) / Avg Earning Assets x 100
- ROA% = (H1 PAT x 2) / Avg Total Assets x 100
- ROE% = (H1 PAT x 2) / Avg Net Worth x 100
- Debt/Equity = Borrowings / Net Worth
- EPS = PAT / Shares Outstanding
Unit conversion: Lakhs÷100=Crores, Millions÷10=Crores
NEVER output N/A. All figures in Rs. Crores.

Respond ONLY with valid JSON:
{{"companyProfile":"...","financialPeriods":["H1FY26","H1FY25","31.03.2025","31.03.2024"],"financialData":[{{"metric":"Total Income","values":["...","...","...","..."]}},{{"metric":"Net Interest Income","values":["...","...","...","..."]}},{{"metric":"Profit After Tax","values":["...","...","...","..."]}},{{"metric":"Net Worth","values":["...","...","...","..."]}},{{"metric":"Total Assets","values":["...","...","...","..."]}},{{"metric":"Loans & Advances","values":["...","...","...","..."]}},{{"metric":"Borrowings","values":["...","...","...","..."]}},{{"metric":"GNPA / Gross Stage 3 (%)","values":["...","...","...","..."]}},{{"metric":"NNPA / Net Stage 3 (%)","values":["...","...","...","..."]}},{{"metric":"CRAR / CAR (%)","values":["...","...","...","..."]}}],"conditionalMetrics":[{{"metric":"Deposits","values":["...","...","...","..."],"includeFor":["bank","apex"]}},{{"metric":"CASA (%)","values":["...","...","...","..."],"includeFor":["bank"]}},{{"metric":"NIM (%)","values":["...","...","...","..."],"includeFor":["bank","nbfc","psu"]}},{{"metric":"ROA (%)","values":["...","...","...","..."],"includeFor":["bank","apex"]}},{{"metric":"ROE (%)","values":["...","...","...","..."],"includeFor":["bank","nbfc","psu"]}},{{"metric":"Tier I CAR (%)","values":["...","...","...","..."],"includeFor":["bank"]}},{{"metric":"Debt/Equity","values":["...","...","...","..."],"includeFor":["nbfc","psu"]}},{{"metric":"EPS (Rs)","values":["...","...","...","..."],"includeFor":["nbfc","psu"]}}],"comments":[{{"heading":"Profitability","text":"..."}},{{"heading":"Asset Quality","text":"..."}},{{"heading":"Capitalisation","text":"..."}},{{"heading":"Liquidity","text":"..."}}],"recommendation":"Keeping in view..."}}"""

            with st.spinner('🤖 Gemini is analyzing and writing your credit report...'):
                try:
                    genai.configure(api_key=api_key)
                    model = genai.GenerativeModel(model_name='gemini-2.5-flash', generation_config=genai.GenerationConfig(temperature=0.1, max_output_tokens=16000))
                    response = model.generate_content(prompt)
                    raw = response.text

                    clean = re.sub(r'```json|```','',raw).strip()
                    match = re.search(r'\{[\s\S]*\}', clean)
                    if match: clean = match.group(0)
                    clean = re.sub(r',\s*}','}',clean); clean = re.sub(r',\s*]',']',clean)

                    try: report_data = json.loads(clean)
                    except:
                        try:
                            import json5; report_data = json5.loads(clean)
                        except: st.error("❌ Parse error — click Generate again"); st.stop()

                    report_data.update({'issuerName':issuer_name,'entityType':entity_code,'preparedBy':prepared_by,'reviewedBy':reviewed_by,'industry':industry,'sector':sector,'reviewPeriod':review_period,'investments':[i for i in investments_data if i.get('security')]})

                    docx_buf = build_docx(report_data)
                    filename = issuer_name.replace(' ','_')+'_Credit_Report.docx'

                    st.success("✅ Report generated!")
                    st.download_button("⬇ Download Word Report", data=docx_buf, file_name=filename, mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)
                    st.markdown("---")
                    st.markdown("### Preview")
                    st.markdown(f"""<div style="background:#16161c;border:1px solid #2a2a35;border-radius:12px;padding:24px">
<p style="color:#c9a84c;font-size:0.7rem;font-weight:700;letter-spacing:1px;text-transform:uppercase">{entity_type}</p>
<h3 style="margin:8px 0 4px;color:#e8e4dc">{issuer_name}</h3>
<p style="color:#9b9688;font-size:0.85rem">{review_period}</p>
<hr style="border-color:#2a2a35;margin:12px 0">
<p style="font-size:0.9rem;color:#c8c4bc;line-height:1.7">{report_data.get('companyProfile','')}</p>
<hr style="border-color:#2a2a35;margin:12px 0">
<p style="font-size:0.85rem;color:#9b9688"><strong style="color:#c9a84c">Recommendation:</strong> {report_data.get('recommendation','')[:250]}...</p>
</div>""", unsafe_allow_html=True)

                except Exception as e:
                    st.error(f"❌ Error: {str(e)}")
